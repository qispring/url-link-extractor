from __future__ import annotations

from typing import List

from openpyxl import Workbook

from .exceptions import ExcelExportError
from .logger import get_logger
from .models import ExtractResult

logger = get_logger("url_link_extractor.excel")


class ExcelExporter:
    def export(self, result: ExtractResult, output_path: str = "result.xlsx") -> str:
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "提取结果"

            ws.cell(row=1, column=1, value="URL")
            ws.cell(row=1, column=2, value="标题")

            for idx, item in enumerate(result.results, start=2):
                ws.cell(row=idx, column=1, value=item.url)
                ws.cell(row=idx, column=2, value=item.title)

            wb.save(output_path)
            logger.info("excel_exported", path=output_path, rows=len(result.results))
            return output_path

        except Exception as exc:
            logger.error("excel_export_failed", path=output_path, error=str(exc))
            raise ExcelExportError(f"Failed to export Excel: {exc}") from exc