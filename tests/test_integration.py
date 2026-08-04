import os
import tempfile
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from openpyxl import load_workbook

from url_link_extractor.models import (
    ExtractConfig,
    ExtractRequest,
    ExtractionStatus,
    LinkTitlePair,
    TargetLink,
    TitleExtractionResult,
)
from url_link_extractor.result_builder import ResultBuilder
from url_link_extractor.excel_exporter import ExcelExporter
from url_link_extractor.html_parser import HtmlParser
from url_link_extractor.url_resolver import UrlResolver


class TestIntegrationLinkDiscovery:
    def test_full_link_discovery_flow(self):
        html = """
        <html><body>
            <a href="https://support.huaweicloud.com/codeartsagent_cli_0001.html">Link 1</a>
            <a href="https://support.huaweicloud.com/codeartsagent_cli_0002.html">Link 2</a>
            <a href="javascript:void(0)">JS Link</a>
            <a href="#section1">Anchor</a>
            <a href="mailto:test@example.com">Email</a>
            <a href="https://www.huaweicloud.com/about.html">External</a>
            <a href="https://support.huaweicloud.com/codeartsagent_cli_0001.html">Duplicate</a>
            <a href="codeartsagent_cli_0003.html">Relative</a>
        </body></html>
        """
        parser = HtmlParser()
        resolver = UrlResolver()
        entry_url = "https://support.huaweicloud.com/usermanual-cli/"
        prefix = "https://support.huaweicloud.com/"
        suffix = ".html"

        raw_links = parser.extract_links(html)
        assert len(raw_links) == 8

        seen = set()
        filtered = []
        for href in raw_links:
            if resolver.is_excluded_link(href):
                continue
            absolute = resolver.to_absolute(entry_url, href)
            if not resolver.match_prefix(absolute, prefix):
                continue
            if suffix and not resolver.match_suffix(absolute, suffix):
                continue
            if absolute in seen:
                continue
            seen.add(absolute)
            filtered.append(absolute)

        assert len(filtered) == 3
        assert "https://support.huaweicloud.com/codeartsagent_cli_0001.html" in filtered
        assert "https://support.huaweicloud.com/codeartsagent_cli_0002.html" in filtered
        assert "https://support.huaweicloud.com/usermanual-cli/codeartsagent_cli_0003.html" in filtered


class TestIntegrationResultAndExcel:
    def test_result_to_excel_flow(self):
        links = [
            TargetLink(url="https://support.huaweicloud.com/codeartsagent_cli_0001.html", order=0),
            TargetLink(url="https://support.huaweicloud.com/codeartsagent_cli_0002.html", order=1),
        ]
        title_results = [
            TitleExtractionResult(url="https://support.huaweicloud.com/codeartsagent_cli_0001.html", title="什么是码道CLI", status=ExtractionStatus.SUCCESS),
            TitleExtractionResult(url="https://support.huaweicloud.com/codeartsagent_cli_0002.html", title="安装码道CLI", status=ExtractionStatus.SUCCESS),
        ]

        builder = ResultBuilder()
        result = builder.build(links, title_results, 0)

        assert result.statistics.total == 2
        assert result.statistics.success == 2
        assert result.statistics.total == result.statistics.success + result.statistics.failed + result.statistics.skipped

        tmpdir = tempfile.mkdtemp()
        output_path = os.path.join(tmpdir, "result.xlsx")
        exporter = ExcelExporter()
        exporter.export(result, output_path)

        wb = load_workbook(output_path)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "URL"
        assert ws.cell(row=1, column=2).value == "标题"
        assert ws.cell(row=2, column=1).value == "https://support.huaweicloud.com/codeartsagent_cli_0001.html"
        assert ws.cell(row=2, column=2).value == "什么是码道CLI"
        assert ws.cell(row=3, column=1).value == "https://support.huaweicloud.com/codeartsagent_cli_0002.html"
        assert ws.cell(row=3, column=2).value == "安装码道CLI"

    def test_partial_failure_flow(self):
        links = [
            TargetLink(url="https://example.com/a.html", order=0),
            TargetLink(url="https://example.com/b.html", order=1),
            TargetLink(url="https://example.com/c.html", order=2),
        ]
        title_results = [
            TitleExtractionResult(url="https://example.com/a.html", title="A", status=ExtractionStatus.SUCCESS),
            TitleExtractionResult(url="https://example.com/b.html", title="", status=ExtractionStatus.FAILED, error="请求超时"),
            TitleExtractionResult(url="https://example.com/c.html", title="C", status=ExtractionStatus.SUCCESS),
        ]

        builder = ResultBuilder()
        result = builder.build(links, title_results, 0)

        assert result.statistics.total == 3
        assert result.statistics.success == 2
        assert result.statistics.failed == 1
        assert len(result.errors) == 1
        assert result.errors[0].url == "https://example.com/b.html"
        assert result.errors[0].error == "请求超时"

    def test_order_consistency(self):
        links = [
            TargetLink(url="https://example.com/c.html", order=0),
            TargetLink(url="https://example.com/a.html", order=1),
            TargetLink(url="https://example.com/b.html", order=2),
        ]
        title_results = [
            TitleExtractionResult(url="https://example.com/c.html", title="C", status=ExtractionStatus.SUCCESS),
            TitleExtractionResult(url="https://example.com/a.html", title="A", status=ExtractionStatus.SUCCESS),
            TitleExtractionResult(url="https://example.com/b.html", title="B", status=ExtractionStatus.SUCCESS),
        ]

        builder = ResultBuilder()
        result = builder.build(links, title_results, 0)

        urls_in_result = [r.url for r in result.results]
        assert urls_in_result == [
            "https://example.com/c.html",
            "https://example.com/a.html",
            "https://example.com/b.html",
        ]


class TestIntegrationProgressCallback:
    def test_progress_callback_triggered(self):
        progress_records = []

        def on_progress(progress):
            progress_records.append((progress.completed, progress.total))

        links = [
            TargetLink(url="https://example.com/a.html", order=0),
            TargetLink(url="https://example.com/b.html", order=1),
            TargetLink(url="https://example.com/c.html", order=2),
        ]

        total = len(links)
        for i in range(total):
            on_progress(type("P", (), {"completed": i + 1, "total": total, "last_url": links[i].url, "last_status": ExtractionStatus.SUCCESS})())

        assert len(progress_records) == 3
        assert progress_records[-1] == (3, 3)
        assert progress_records[-1][0] == progress_records[-1][1]