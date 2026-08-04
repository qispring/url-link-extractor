import os
import tempfile

import pytest
from openpyxl import load_workbook

from url_link_extractor.excel_exporter import ExcelExporter
from url_link_extractor.exceptions import ExcelExportError
from url_link_extractor.models import (
    ExtractResult,
    ExtractStatistics,
    ExtractionStatus,
    LinkTitlePair,
)


class TestExcelExporter:
    def setup_method(self):
        self.exporter = ExcelExporter()
        self._tmpdir = tempfile.mkdtemp()

    def _make_result(self):
        results = [
            LinkTitlePair(url="https://example.com/a.html", title="什么是码道CLI", status=ExtractionStatus.SUCCESS),
            LinkTitlePair(url="https://example.com/b.html", title="测试标题", status=ExtractionStatus.SUCCESS),
        ]
        stats = ExtractStatistics(total=2, success=2, failed=0, skipped=0, duration=1000)
        return ExtractResult(results=results, statistics=stats, errors=[])

    def test_export_basic(self):
        output_path = os.path.join(self._tmpdir, "result.xlsx")
        result = self._make_result()
        path = self.exporter.export(result, output_path)

        assert os.path.exists(path)
        wb = load_workbook(path)
        ws = wb.active

        assert ws.cell(row=1, column=1).value == "URL"
        assert ws.cell(row=1, column=2).value == "标题"
        assert ws.cell(row=2, column=1).value == "https://example.com/a.html"
        assert ws.cell(row=2, column=2).value == "什么是码道CLI"
        assert ws.cell(row=3, column=1).value == "https://example.com/b.html"
        assert ws.cell(row=3, column=2).value == "测试标题"

    def test_export_chinese_title(self):
        output_path = os.path.join(self._tmpdir, "chinese.xlsx")
        result = self._make_result()
        path = self.exporter.export(result, output_path)

        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(row=2, column=2).value == "什么是码道CLI"

    def test_export_empty_result(self):
        output_path = os.path.join(self._tmpdir, "empty.xlsx")
        result = ExtractResult(
            results=[],
            statistics=ExtractStatistics(0, 0, 0, 0, 0),
            errors=[],
        )
        path = self.exporter.export(result, output_path)

        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "URL"
        assert ws.cell(row=1, column=2).value == "标题"

    def test_export_order_preserved(self):
        output_path = os.path.join(self._tmpdir, "order.xlsx")
        results = [
            LinkTitlePair(url="https://example.com/c.html", title="C", status=ExtractionStatus.SUCCESS),
            LinkTitlePair(url="https://example.com/a.html", title="A", status=ExtractionStatus.SUCCESS),
            LinkTitlePair(url="https://example.com/b.html", title="B", status=ExtractionStatus.SUCCESS),
        ]
        result = ExtractResult(
            results=results,
            statistics=ExtractStatistics(3, 3, 0, 0, 0),
            errors=[],
        )
        path = self.exporter.export(result, output_path)

        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "https://example.com/c.html"
        assert ws.cell(row=3, column=1).value == "https://example.com/a.html"
        assert ws.cell(row=4, column=1).value == "https://example.com/b.html"